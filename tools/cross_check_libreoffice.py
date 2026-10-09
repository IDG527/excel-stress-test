"""Check the reference model against the workbook's real formulas, without Excel.

Fills the 60 form rows of a COPY of the workbook with random entries, lets
LibreOffice recalculate it, and compares every row with tools/oracle.py.
The Power Query is not refreshed or edited; its last loaded holiday table is used.

    pip install openpyxl
    sudo apt-get install libreoffice-calc
    python tools/cross_check_libreoffice.py "OVERTIME_FORM_2026_MM__NAME__rev_1.2_QUERY_LINKED.xlsx" --runs 9

Exit code 0 = every row matched.
"""
import argparse
import datetime as dt
import os
import random
import shutil
import subprocess
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from oracle import oracle  # noqa: E402

FORM, HOLS = "Engineer Name", "SG Public Holidays"
MASKS = [(1, 1, 1, 1)] * 3 + [(1, 1, 0, 0), (0, 0, 1, 1), (1, 0, 0, 1), (1, 1, 0, 1), (1, 0, 1, 1),
                              (0, 1, 1, 1), (1, 1, 1, 0), (1, 0, 0, 0), (0, 0, 0, 0), (1, 0, 1, 0), (0, 1, 0, 1)]


def random_rows(rng, holidays, year):
    rows = []
    for _ in range(60):
        r = rng.random()
        if r < 0.05:
            date = None
        elif r < 0.30:
            date = rng.choice(holidays)
        else:
            date = dt.date(year, 1, 1) + dt.timedelta(days=rng.randrange(365))
        mask = rng.choice(MASKS)
        snap = rng.random() < 0.5
        start = rng.randrange(96) * 15 if snap else rng.randrange(1440)
        span = rng.randrange(1439)                      # < 24 h: at most one midnight roll-over
        pts = sorted(rng.randrange(span + 1) for _ in range(sum(mask)))
        if snap:
            pts = sorted(p // 15 * 15 for p in pts)
        it = iter((start + p) % 1440 for p in pts)
        rows.append((date, [next(it) if m else None for m in mask]))
    return rows


def recalc(path, outdir):
    subprocess.run(["soffice", "--headless", "--norestore", "--calc", "--convert-to", "xlsx",
                    "--outdir", outdir, path], check=True, capture_output=True, timeout=600)
    return os.path.join(outdir, os.path.basename(path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--runs", type=int, default=3)
    ap.add_argument("--seed", type=int, default=20260929)
    ap.add_argument("--office-hours", metavar="HH:MM-HH:MM",
                    help="put these office hours in Q1/Q2 of the copy first, e.g. 09:00-18:00")
    args = ap.parse_args()

    rng = random.Random(args.seed)
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "copy.xlsx")
    shutil.copyfile(args.workbook, src)
    wb = openpyxl.load_workbook(src)
    hol_ws = wb[HOLS]
    holidays = [c.value.date() for c in hol_ws["A"][1:] if isinstance(c.value, dt.datetime)]
    if args.office_hours:                       # the copy only; openpyxl ignores sheet protection
        start, end = (dt.datetime.strptime(t, "%H:%M").time() for t in args.office_hours.split("-"))
        wb[FORM]["Q1"].value, wb[FORM]["Q2"].value = start, end
        wb.save(src)
    q1 = wb[FORM]["Q1"].value
    q2 = wb[FORM]["Q2"].value
    q1, q2 = q1.hour * 60 + q1.minute, q2.hour * 60 + q2.minute
    print(f"office hours {q1 // 60:02d}:{q1 % 60:02d}-{q2 // 60:02d}:{q2 % 60:02d}")
    year = holidays[0].year

    bad_total = 0
    for run in range(1, args.runs + 1):
        wb = openpyxl.load_workbook(src)
        ws = wb[FORM]
        rows = random_rows(rng, holidays, year)
        for i, (date, times) in enumerate(rows):
            r = 11 + i
            for col in "ADEFGJKL":
                ws[f"{col}{r}"].value = None
            if date:
                ws[f"A{r}"].value = dt.datetime.combine(date, dt.time())
            for col, t in zip("DEFG", times):
                if t is not None:
                    ws[f"{col}{r}"].value = dt.time(t // 60, t % 60)
        filled = os.path.join(tmp, f"run{run}.xlsx")
        wb.save(filled)
        out = openpyxl.load_workbook(recalc(filled, os.path.join(tmp, "out")), data_only=True)[FORM]

        bad = 0
        for i, (date, times) in enumerate(rows):
            r = 11 + i
            exp = oracle(date, times, set(holidays), q1, q2)
            got = tuple("" if out.cell(r, c).value is None else out.cell(r, c).value for c in (2, 3, 8, 9))
            same_day = got[0] == exp[0] or (date is None and str(got[0]).strip() == "")
            same_num = all(a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-3)
                           for a, b in zip(got[2:], exp[2:]))
            if not (same_day and got[1] == exp[1] and same_num):
                bad += 1
                print(f"run {run} row {r}: {date} {times} form={got} model={exp}")
        print(f"run {run}: 60 rows, {bad} mismatches, totals travel={out['H71'].value} work={out['I71'].value}")
        bad_total += bad

    shutil.rmtree(tmp, ignore_errors=True)
    print("ALL ROWS MATCH" if bad_total == 0 else f"{bad_total} MISMATCHES")
    sys.exit(1 if bad_total else 0)


if __name__ == "__main__":
    main()
